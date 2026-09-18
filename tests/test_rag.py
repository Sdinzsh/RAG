import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import requests
import rag_engine as rag
import app as web


def make_tree(sections, source_format="pdf"):
    rag._link_parents(sections)
    tree = rag.DocumentTree("fixture.pdf", 4, sections, source_format=source_format)
    tree.build_index()
    return tree


def section(sid, title, level=1, text="evidence"):
    return rag.Section(sid, title, level, 1, 1, text)


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        cache = patch.object(rag, "CACHE_DIR", self.root / "cache")
        cache.start()
        self.addCleanup(cache.stop)

    def build(self, text, name="sample.md"):
        path = self.root / name
        path.write_text(text)
        return rag.build_document_tree(str(path), name)

    def test_markdown_depth_and_closed_fences(self):
        tree = self.build("# A\nintro\n## B\n```python\n# fake\n```\n### C\nbody\n#### D\nleaf")
        self.assertEqual([s.title for s in tree.sections], ["A", "B", "C", "D"])
        self.assertEqual(tree.sections[-1].parent_id, tree.sections[-2].section_id)
        self.assertIn("# fake", tree.sections[1].text)
        self.assertEqual(tree.sections[-1].location, "Block 4")

    def test_setext_heading_not_duplicated_in_previous_body(self):
        tree = self.build("Introduction\n===\nhello\n\nDetails\n---\nworld")
        self.assertEqual([s.title for s in tree.sections], ["Introduction", "Details"])
        self.assertNotIn("Details", tree.sections[0].text)
        self.assertEqual(tree.sections[1].parent_id, tree.sections[0].section_id)

    def test_pdf_two_headings_same_page_and_continuation(self):
        import pymupdf
        path = self.root / "report.pdf"
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.insert_text((72, 72), "Subscriptions", fontsize=18)
            page.insert_text((72, 100), "Premium monthly income and offline downloads.", fontsize=11)
            page.insert_text((72, 160), "Advertising", fontsize=18)
            page.insert_text((72, 190), "Free listeners hear targeted audio ads.", fontsize=11)
            page = doc.new_page()
            page.insert_text((72, 72), "Campaign analytics support advertisers.", fontsize=11)
            doc.save(path)
        tree = rag.build_document_tree(str(path), path.name)
        self.assertEqual([s.title for s in tree.sections], ["Subscriptions", "Advertising"])
        self.assertNotIn("Advertising", tree.sections[0].text)
        self.assertEqual(tree.sections[0].location, "Page 1")
        self.assertEqual(tree.sections[1].location, "Pages 1-2")
        self.assertIn("Campaign analytics", tree.sections[1].text)

    def test_docx_heading_depth_and_table(self):
        from docx import Document
        doc = Document()
        doc.add_heading("Revenue", 1)
        doc.add_heading("Premium", 2)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Price"
        table.cell(0, 1).text = "15 credits"
        path = self.root / "sample.docx"
        doc.save(path)
        tree = rag.build_document_tree(str(path), path.name)
        self.assertEqual(tree.sections[1].parent_id, tree.sections[0].section_id)
        self.assertIn("15 credits", tree.sections[1].text)
        self.assertEqual(tree.sections[1].location, "Block 2")

    def test_empty_document_rejected(self):
        with self.assertRaisesRegex(ValueError, "No readable text"):
            self.build(" \n\n")

    def test_headingless_document_has_flat_blocks(self):
        tree = self.build("First paragraph.\n\nSecond paragraph.", "sample.txt")
        self.assertEqual(len(tree.roots), 2)
        self.assertEqual(tree.sections[1].location, "Block 2")

    def test_cache_reused_versioned_and_format_aware(self):
        first = self.build("# Title\nbody")
        second = self.build("# Title\nbody", "renamed.md")
        self.assertEqual(first.doc_id, second.doc_id)
        self.assertTrue(second.from_cache)
        self.assertEqual(second.filename, "renamed.md")
        self.assertNotEqual(first.doc_id, self.build("# Title\nbody", "sample.txt").doc_id)
        cache = rag._cache_path(first.doc_id)
        content = json.loads(cache.read_text())
        content["cache_version"] = -1
        cache.write_text(json.dumps(content))
        self.assertIsNone(rag.load_cached_tree(first.doc_id))
        self.assertFalse(self.build("# Title\nbody").from_cache)

    def test_invalid_cache_id(self):
        self.assertIsNone(rag.load_cached_tree("../README"))


class RetrievalTests(unittest.TestCase):
    def test_all_nodes_seen_beyond_thirty_and_deep_parent_content(self):
        sections = [section(f"s_{i:03}", f"Root {i}") for i in range(35)]
        sections += [section("child", "Child", 2), section("grandchild", "Grandchild", 3)]
        tree = make_tree(sections)
        seen = set()
        targets = {"s_034", "grandchild"}

        def search(model, messages, response_format):
            entries = json.loads(messages[-1]["content"])["tree"]
            seen.update(entry["id"] for entry in entries)
            for entry in entries:
                if entry["id"] == "grandchild":
                    self.assertEqual(entry["ancestors"], ["Root 34", "Child"])
            return json.dumps({"node_list": [e["id"] for e in entries if e["id"] in targets]})

        with patch.object(rag, "ollama_chat", side_effect=search):
            found = rag.find_relevant_sections(tree, "question", "local")
        self.assertEqual(seen, {s.section_id for s in sections})
        self.assertEqual({s.section_id for s in found}, targets)

    def test_large_selection_reduced_to_top_k(self):
        tree = make_tree([section(str(i), f"Section {i}") for i in range(100)])

        def search(model, messages, response_format):
            ids = [s["id"] for s in json.loads(messages[-1]["content"])["tree"]]
            return json.dumps({"node_list": ids[:4]})

        with patch.object(rag, "ollama_chat", side_effect=search) as chat:
            result = rag.find_relevant_sections(tree, "question", "local")
        self.assertEqual(len(result), 4)
        self.assertGreater(chat.call_count, 7)

    def test_valid_no_match_skips_generation(self):
        tree = make_tree([section("a", "Revenue")])
        with patch.object(rag, "ollama_chat", return_value='{"node_list": []}'), patch.object(rag, "ollama_chat_stream") as generate:
            events = list(rag.answer_query_stream(tree, "unknown fact", "local"))
        generate.assert_not_called()
        self.assertEqual(events[0], ("sections", []))
        self.assertIn("could not find", events[1][1])
        self.assertEqual(events[-1][0], "done")

    def test_invalid_json_and_unknown_ids_use_keyword_fallback(self):
        tree = make_tree([section("a", "Revenue", text="subscription"), section("b", "Users")])
        for raw in ['not json', '{"node_list": ["invented"]}', '{"node_list": [4]}']:
            with self.subTest(raw=raw), patch.object(rag, "ollama_chat", return_value=raw):
                self.assertEqual(rag.find_relevant_sections(tree, "subscription", "local"), [tree.get("a")])

    def test_fallback_does_not_invent_relevance(self):
        tree = make_tree([section("a", "Revenue")])
        self.assertEqual(rag._keyword_fallback(tree, "penguins", 4), [])

    def test_duplicate_ids_collapsed(self):
        self.assertEqual(rag._extract_id_array('{"node_list": ["a", "a"]}', 4), ["a"])
        self.assertEqual(rag._extract_id_array('[]', 4), [])

    def test_model_failure_propagates(self):
        tree = make_tree([section("a", "Revenue")])
        with patch.object(rag, "ollama_chat", side_effect=requests.ConnectionError("offline")):
            with self.assertRaises(requests.ConnectionError):
                rag.find_relevant_sections(tree, "Revenue", "local")

    def test_full_context_and_explicit_limit(self):
        source = section("a", "Subscriptions", text="x" * 3500 + " evidence at the end")
        self.assertIn("evidence at the end", rag.build_context([source]))
        self.assertIn("[a] Subscriptions (Page 1)", rag.build_context([source]))
        with patch.object(rag, "MAX_CONTEXT_CHARS", 50), self.assertRaisesRegex(ValueError, "context budget"):
            rag.build_context([source])

    def test_followup_uses_previous_user_question(self):
        tree = make_tree([section("a", "Revenue")])
        with patch.object(rag, "find_relevant_sections", return_value=[]) as find:
            list(rag.answer_query_stream(tree, "How does it work?", "local", [
                {"role": "user", "content": "Describe advertising"},
                {"role": "assistant", "content": "untrusted previous answer"},
            ]))
        self.assertIn("Describe advertising", find.call_args.args[1])
        self.assertNotIn("untrusted previous answer", find.call_args.args[1])

    def test_new_topic_does_not_inherit_old_question_or_answer(self):
        tree = make_tree([section("spotify", "Spotify architecture", text="Spotify-like app uses Next.js.")])
        history = [
            {"role": "user", "content": "what is opencode"},
            {"role": "assistant", "content": 'The term "opencode" does not appear in the provided document sections.'},
        ]
        for question in ["spotify", "What is Spotify?", "What about Spotify?"]:
            with self.subTest(question=question):
                with patch.object(rag, "find_relevant_sections", return_value=tree.sections) as find, patch.object(rag, "ollama_chat_stream", return_value=iter(["A Spotify-like app uses Next.js."])) as generate:
                    list(rag.answer_query_stream(tree, question, "local", history))
                self.assertEqual(find.call_args.args[1], question)
                messages = generate.call_args.args[1]
                self.assertEqual([m["role"] for m in messages], ["system", "user"])
                self.assertNotIn("opencode", json.dumps(messages).lower())
                self.assertIn(question, messages[-1]["content"])

    def test_followup_context_stops_at_latest_topic_switch(self):
        tree = make_tree([section("a", "Spotify")])
        history = [
            {"role": "user", "content": "what is opencode"},
            {"role": "assistant", "content": "old refusal"},
            {"role": "user", "content": "spotify"},
            {"role": "assistant", "content": "old generated answer"},
            {"role": "user", "content": "How does it work?"},
        ]
        with patch.object(rag, "find_relevant_sections", return_value=tree.sections) as find, patch.object(rag, "ollama_chat_stream", return_value=iter(["Document evidence."])) as generate:
            list(rag.answer_query_stream(tree, "Tell me more", "local", history))
        search = find.call_args.args[1]
        self.assertIn("spotify", search)
        self.assertIn("How does it work?", search)
        self.assertIn("Tell me more", search)
        self.assertNotIn("opencode", search)
        messages = generate.call_args.args[1]
        self.assertIn(search, messages[-1]["content"])
        self.assertNotIn("old generated answer", json.dumps(messages))

    def test_source_footer_when_model_omits_citations(self):
        tree = make_tree([section("a", "Revenue")])
        with patch.object(rag, "find_relevant_sections", return_value=tree.sections), patch.object(rag, "ollama_chat_stream", return_value=iter(["Subscription income."])):
            events = list(rag.answer_query_stream(tree, "Revenue?", "local"))
        self.assertIn("Retrieved sources:\n- [a] Revenue (Page 1)", events[-2][1])
        self.assertEqual(events[-1][0], "done")

    def test_empty_generation_is_an_error(self):
        tree = make_tree([section("a", "Revenue")])
        with patch.object(rag, "find_relevant_sections", return_value=tree.sections), patch.object(rag, "ollama_chat_stream", return_value=iter([])):
            with self.assertRaisesRegex(RuntimeError, "empty answer"):
                list(rag.answer_query_stream(tree, "Revenue?", "local"))


class OllamaTests(unittest.TestCase):
    def test_cloud_model_rejected_before_network(self):
        with patch.object(rag, "_local_http") as http, self.assertRaises(ValueError):
            list(rag.ollama_chat_stream("model:cloud", []))
        http.assert_not_called()

    def test_aliased_remote_model_rejected(self):
        http = MagicMock()
        http.__enter__.return_value = http
        http.post.return_value.json.return_value = {"remote_model": "remote"}
        with patch.object(rag, "_local_http", return_value=http), self.assertRaisesRegex(ValueError, "remotely"):
            rag.validate_local_model("innocent-alias")

    def test_environment_proxy_disabled(self):
        with rag._local_http() as session:
            self.assertFalse(session.trust_env)

    def test_stream_errors_and_incomplete_streams(self):
        for chunks, error in [([b'{"error":"out of memory"}'], "out of memory"),
                              ([b'{"message":{"content":"partial"}}'], "before completion")]:
            http = MagicMock()
            http.__enter__.return_value = http
            response = http.post.return_value.__enter__.return_value
            response.iter_lines.return_value = chunks
            with patch.object(rag, "validate_local_model"), patch.object(rag, "_local_http", return_value=http):
                with self.assertRaisesRegex(RuntimeError, error):
                    list(rag.ollama_chat_stream("local", []))


class APITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        for name, target, value in [("CACHE_DIR", rag, self.root / "cache"),
                                     ("UPLOAD_FOLDER", web, self.root), ("SESSIONS", web, {})]:
            patcher = patch.object(target, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.client = web.app.test_client()

    def upload(self):
        return self.client.post('/api/upload', data={
            "file": (io.BytesIO(b'# Revenue\n## Premium\nSubscriptions fund the service.'), 'music.md'),
            "model": "local",
        })

    def test_upload_export_resume_and_chat(self):
        response = self.upload()
        self.assertEqual(response.status_code, 200)
        document = response.get_json()
        self.assertEqual(document["location_unit"], "blocks")
        tree = self.client.get(f'/api/doc/{document["doc_id"]}/tree').get_json()
        self.assertEqual(tree["result"][0]["nodes"][0]["title"], "Premium")
        self.assertIsNone(tree["result"][0]["page_index"])
        self.assertTrue(self.client.get(f'/api/doc/{document["doc_id"]}').get_json()["cached"])
        self.assertFalse(list(self.root.glob('tmp*')))
        with patch.object(rag, "ollama_chat", return_value='{"node_list":["s_002"]}'), patch.object(rag, "ollama_chat_stream", return_value=iter(['Subscription income [s_002, Premium, Block 2].'])):
            response = self.client.post('/api/chat', json={"session_id": document["session_id"], "query": "Revenue?"})
            events = [json.loads(line[6:]) for line in response.get_data(as_text=True).splitlines() if line.startswith('data: ')]
        self.assertEqual([e["type"] for e in events], ["sections", "token", "done"])
        self.assertEqual(events[0]["data"][0]["title"], "Premium")
        self.assertEqual(web.SESSIONS[document["session_id"]]["history"][-1]["role"], "assistant")

    def test_failed_answer_not_saved_to_history(self):
        session = self.upload().get_json()["session_id"]
        with patch.object(rag, "find_relevant_sections", side_effect=RuntimeError("model unavailable")):
            response = self.client.post('/api/chat', json={"session_id": session, "query": "Revenue?"})
            self.assertIn('"type": "error"', response.get_data(as_text=True))
        self.assertEqual(web.SESSIONS[session]["history"], [])

    def test_pdf_chat_retrieves_only_the_uploaded_sessions_pdf(self):
        import pymupdf

        def upload_pdf(name, fact):
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 72), "Delivery policy", fontsize=18)
                page.insert_text((72, 100), fact, fontsize=11)
                content = pdf.tobytes()
            response = self.client.post('/api/upload', data={
                "file": (io.BytesIO(content), name), "model": "local",
            })
            self.assertEqual(response.status_code, 200)
            return response.get_json()

        first = upload_pdf("north.pdf", "North delivery takes 17 days.")
        second = upload_pdf("south.pdf", "South delivery takes 29 days.")
        self.assertNotEqual(first["doc_id"], second["doc_id"])
        for document, expected, excluded in [
            (first, "17 days", "29 days"), (second, "29 days", "17 days")
        ]:
            with self.subTest(filename=document["filename"]):
                with patch.object(rag, "ollama_chat", return_value='{"node_list":["s_001"]}'), patch.object(rag, "ollama_chat_stream", return_value=iter([expected])) as generate:
                    response = self.client.post('/api/chat', json={
                        "session_id": document["session_id"], "query": "How long does delivery take?",
                    })
                    events = [json.loads(line[6:]) for line in response.get_data(as_text=True).splitlines() if line.startswith('data: ')]
                context = generate.call_args.args[1][-1]["content"]
                self.assertIn(document["filename"], context)
                self.assertIn(expected, context)
                self.assertNotIn(excluded, context)
                self.assertEqual(events[0]["data"][0]["location"], "Page 1")
                self.assertIn("Delivery policy (Page 1)", events[-2]["data"])
                self.assertEqual(events[-1]["type"], "done")

    def test_page_only_uses_local_assets(self):
        html = self.client.get('/').get_data(as_text=True)
        self.assertNotIn('https://', html)
        self.assertIn('/static/markdown.js', html)
        with self.client.get('/static/markdown.js') as response:
            self.assertEqual(response.status_code, 200)

    def test_bad_chat_input(self):
        for body in [[1], {"query": [1], "session_id": "none"}, {"session_id": []}]:
            self.assertEqual(self.client.post('/api/chat', json=body).status_code, 400)


if __name__ == '__main__':
    unittest.main()
